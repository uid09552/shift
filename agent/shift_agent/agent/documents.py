"""
Turning an uploaded roster file into a grid the agent can read.

A ward's existing shift plan arrives as whatever the ward already uses: a
spreadsheet, a CSV export, or a PDF printed from one. This module reduces all
three to the same shape — a list of sheets, each a rectangular grid of trimmed
cell strings — so the rest of the upload path (roster.py) never has to care
which format it came from.

Nothing here interprets the content. Which column holds the names, whether the
dates run across the top or down the side, what "N" means — that is the model's
job, working from the preview this produces and declaring the layout back to
``interpretRosterUpload``. Deciding it here would mean guessing at parse time,
with no way for the user to correct the guess.

The PDF path is the weakest of the three, as it must be: a PDF has no cells,
only glyphs at coordinates. ``pdfplumber`` recovers ruled tables well and
whitespace-aligned columns tolerably; a roster printed without either comes
through as text lines, and the model is told so.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)

# What the chat upload endpoint accepts, by extension.
SUPPORTED_EXTENSIONS = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".pdf")

# Guardrails on what one upload may expand to. A roster is at most a few
# hundred rows over ~40 day columns; anything far past that is a file that was
# never a roster, and parsing it in full only burns memory and context.
MAX_ROWS_PER_SHEET = 2000
MAX_COLS_PER_ROW = 200
MAX_SHEETS = 12


class DocumentError(Exception):
    """The upload could not be read as a table — reported to the user as-is."""


@dataclass
class Sheet:
    """One rectangular grid of cell strings, with its origin named."""

    name: str
    rows: list[list[str]] = field(default_factory=list)
    # Set when the rows were recovered from prose rather than real cells (a PDF
    # with no ruled table), so the preview can warn that the columns are a
    # best-effort split.
    approximate: bool = False

    @property
    def height(self) -> int:
        return len(self.rows)

    @property
    def width(self) -> int:
        return max((len(r) for r in self.rows), default=0)


@dataclass
class ParsedDocument:
    filename: str
    kind: str
    sheets: list[Sheet] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not any(sheet.rows for sheet in self.sheets)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _clean(value: object) -> str:
    """A cell as a single-line, trimmed string. None and NaN become empty."""
    if value is None:
        return ""
    if isinstance(value, float) and value != value:  # NaN
        return ""
    text = str(value).strip()
    # Excel and PDF cells can carry embedded newlines; a grid cell is one line.
    return " ".join(text.split())


def _trim(rows: list[list[str]], *, drop_blank: bool = False) -> list[list[str]]:
    """Drop wholly empty leading/trailing rows and cap the grid's size.

    Interior blank rows are kept by default: in a spreadsheet they usually
    separate teams or precede a subheading, and that is information. They are
    dropped only where they are a parsing artefact rather than content — see
    the PDF text-strategy path, which emits one between every real row.
    """
    rows = [r[:MAX_COLS_PER_ROW] for r in rows[:MAX_ROWS_PER_SHEET]]
    if drop_blank:
        rows = [r for r in rows if any(r)]
    while rows and not any(rows[0]):
        rows.pop(0)
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


# ---------------------------------------------------------------------------
# Per-format parsers
# ---------------------------------------------------------------------------


def _decode(data: bytes) -> str:
    """Text from bytes, trying the encodings spreadsheet exports actually use.

    utf-8-sig first so a BOM doesn't end up glued to the first header cell,
    then cp1252, which is what Excel writes on a German Windows machine and
    which would otherwise mangle every umlaut in the name column.
    """
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_csv(filename: str, data: bytes) -> ParsedDocument:
    text = _decode(data)
    if not text.strip():
        raise DocumentError("The file is empty.")

    # Sniff over a sample rather than the whole file — Sniffer is quadratic-ish
    # and a month-long roster is long. Semicolon is listed explicitly because
    # it is the default in German Excel, where Sniffer often picks the comma
    # out of the decimal numbers instead.
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:8192], delimiters=",;\t|"
        )
    except csv.Error:
        dialect = csv.excel

    rows = [[_clean(cell) for cell in row] for row in csv.reader(io.StringIO(text), dialect)]
    return ParsedDocument(
        filename=filename,
        kind="csv",
        sheets=[Sheet(name=filename, rows=_trim(rows))],
    )


def _parse_xlsx(filename: str, data: bytes) -> ParsedDocument:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentError(
            "Spreadsheet support is not installed on the server (openpyxl)."
        ) from exc

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentError("Could not read the file as an Excel workbook.") from exc

    sheets: list[Sheet] = []
    try:
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            rows = [
                [_clean(cell) for cell in row]
                for row in worksheet.iter_rows(max_row=MAX_ROWS_PER_SHEET, values_only=True)
            ]
            trimmed = _trim(rows)
            if trimmed:
                sheets.append(Sheet(name=worksheet.title, rows=trimmed))
    finally:
        workbook.close()

    if not sheets:
        raise DocumentError("The workbook has no non-empty sheets.")
    return ParsedDocument(filename=filename, kind="xlsx", sheets=sheets)


# Two or more spaces: what separates columns in a PDF rendered without any
# recoverable structure at all, the last of the three fallbacks below.
_COLUMN_GAP = re.compile(r"\s{2,}")

# pdfplumber settings that infer a grid from where the text sits rather than
# from ruled lines. This is what rescues the common case: a roster printed as a
# plain aligned table, no borders. Column boundaries come out of the word
# positions, so the cells are real cells — just not guaranteed ones.
_TEXT_GRID = {"vertical_strategy": "text", "horizontal_strategy": "text"}


def _parse_pdf(filename: str, data: bytes) -> ParsedDocument:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentError("PDF support is not installed on the server (pdfplumber).") from exc

    sheets: list[Sheet] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                for sheet in _pdf_page_sheets(page, page_no):
                    sheets.append(sheet)
                    if len(sheets) >= MAX_SHEETS:
                        break
                if len(sheets) >= MAX_SHEETS:
                    break
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("Could not read the file as a PDF.") from exc

    if not sheets:
        raise DocumentError(
            "No text could be extracted from the PDF — it is probably a scan. "
            "A CSV or Excel export of the same plan would work."
        )
    return ParsedDocument(filename=filename, kind="pdf", sheets=sheets)


def _pdf_page_sheets(page: object, page_no: int) -> list[Sheet]:
    """One PDF page as sheets, trying three ways of finding its columns.

    In descending order of trustworthiness: ruled table borders, text
    alignment, then plain lines split on wide gaps. Only the first gives cell
    boundaries the document itself asserts, so the other two are flagged
    ``approximate`` and the preview says so.
    """
    # 1. Ruled tables — real cell boundaries drawn in the document.
    tables = page.extract_tables() or []  # type: ignore[attr-defined]
    if tables:
        sheets = []
        for table_no, table in enumerate(tables, start=1):
            rows = _trim([[_clean(cell) for cell in row] for row in table])
            if rows:
                suffix = f" table {table_no}" if len(tables) > 1 else ""
                sheets.append(Sheet(name=f"page {page_no}{suffix}", rows=rows))
        if sheets:
            return sheets

    # 2. Columns inferred from where the words sit. A blank row between every
    # real one is this strategy's signature, hence drop_blank.
    try:
        inferred = page.extract_tables(_TEXT_GRID) or []  # type: ignore[attr-defined]
    except Exception:
        inferred = []
    for table in inferred:
        rows = _trim([[_clean(cell) for cell in row] for row in table], drop_blank=True)
        # One column means it found no columns at all — no better than the
        # text fallback below, and misleading to present as a grid.
        if rows and max(len(r) for r in rows) > 1:
            return [Sheet(name=f"page {page_no}", rows=rows, approximate=True)]

    # 3. Text lines split on wide gaps. layout=True keeps the horizontal
    # spacing that makes those gaps meaningful.
    text = page.extract_text(layout=True) or ""  # type: ignore[attr-defined]
    rows = _trim(
        [_COLUMN_GAP.split(line.strip()) for line in text.splitlines() if line.strip()],
        drop_blank=True,
    )
    if rows:
        return [Sheet(name=f"page {page_no}", rows=rows, approximate=True)]
    return []


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    """Parse an uploaded file into sheets of cell strings.

    Args:
        filename: the client-supplied name; only its extension is trusted, and
            only to pick a parser.
        data: the raw bytes.

    Raises:
        DocumentError: the extension is unsupported, or the bytes could not be
            read as that format. The message is written for the end user.
    """
    lower = filename.lower()
    if lower.endswith((".csv", ".tsv", ".txt")):
        parsed = _parse_csv(filename, data)
    elif lower.endswith((".xlsx", ".xlsm")):
        parsed = _parse_xlsx(filename, data)
    elif lower.endswith(".pdf"):
        parsed = _parse_pdf(filename, data)
    elif lower.endswith(".xls"):
        # The legacy binary format openpyxl cannot read. Worth naming
        # explicitly — "unsupported file type" would look like a bug to
        # someone whose ward still saves .xls.
        raise DocumentError(
            "The old .xls format cannot be read. Save the file as .xlsx or CSV and try again."
        )
    else:
        raise DocumentError(
            f"Unsupported file type. Upload one of: {', '.join(SUPPORTED_EXTENSIONS)}."
        )

    if parsed.is_empty:
        raise DocumentError("The file contains no readable rows.")
    return parsed


def render_grid(
    sheet: Sheet,
    *,
    first_row: int = 0,
    max_rows: int = 30,
    max_cols: int = 40,
) -> str:
    """A slice of a sheet as an indexed text grid for the model to read.

    Both axes are labelled with their indices, because those indices are the
    vocabulary the model answers in: ``interpretRosterUpload`` is told which
    *column number* holds the names and which *row number* holds the dates, and
    it can only give a number it has seen.
    """
    rows = sheet.rows[first_row : first_row + max_rows]
    if not rows:
        return "(no rows in this range)"

    width = min(max(len(r) for r in rows), max_cols)
    header = "row | " + " | ".join(f"c{c}" for c in range(width))
    lines = [header, "-" * len(header)]
    for offset, row in enumerate(rows):
        cells = [(row[c] if c < len(row) else "") for c in range(width)]
        lines.append(f"{first_row + offset:>3} | " + " | ".join(cells))

    shown = first_row + len(rows)
    if shown < sheet.height:
        lines.append(f"… {sheet.height - shown} more row(s) not shown (rows {shown}–{sheet.height - 1})")
    if sheet.width > width:
        lines.append(f"… {sheet.width - width} more column(s) not shown")
    return "\n".join(lines)
